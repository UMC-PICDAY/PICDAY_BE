# 목업 엑셀(v7) -> prisma/seed-data/studios.json 생성기
#
# 시드를 돌릴 때는 필요 없다. 이미 생성된 studios.json이 커밋돼 있으므로
# `pnpm seed:studios` 만 실행하면 된다.
# 이 스크립트는 **목업 엑셀이 갱신됐을 때 JSON을 다시 만드는 용도**다.
#
#   준비:  pip install openpyxl
#   실행:  python prisma/seed-data/generate-from-xlsx.py <목업엑셀.xlsx>
#
# 목업 엑셀은 용량·보안 문제로 레포에 두지 않는다. 노션/드라이브에서 받아
# 경로를 인자로 넘길 것. 이미지 매핑(image-mapping.csv)은 이 폴더에 함께 있다.

import openpyxl, re, json, csv, os, sys, collections

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

XLSX = sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE_DIR, "PICDAY_목업데이터_v7.xlsx")
IMG_CSV = os.path.join(BASE_DIR, "image-mapping.csv")
OUT = os.path.join(BASE_DIR, "studios.json")
# 엑셀 버전에 따라 시트명 접미사가 붙기도 한다 (v7: "..._v12", v8: 접미사 없음)
SHEET_CANDIDATES = ["사진관 목업 데이터", "사진관 목업 데이터_v12"]

if not os.path.exists(XLSX):
    sys.exit(
        f"목업 엑셀을 찾을 수 없습니다: {XLSX}\n"
        f"사용법: python {os.path.basename(__file__)} <목업엑셀.xlsx>"
    )

LOCATION = {
    "홍대": "HONGDAE", "강남": "GANGNAM", "성수": "SEONGSU", "연남": "YEONNAM",
    "건대": "KONDAE", "신촌": "SINCHON", "잠실": "JAMSIL", "압구정": "APGUJEONG",
    "혜화": "HYEHWA", "종로": "JONGNO",
}

# 1~9: n호선 / 11: 공항철도 / 12: 경의중앙선 / 13: 신분당선 (API 명세 기준)
LINE_CODE = {"공항철도": 11, "경의중앙선": 12, "신분당선": 13}

# 상품명 접두어 -> (ShootingCategory, 기본 인원)
CATEGORY = [
    ("증명", "ID_PHOTO", 1),
    ("여권", "ID_PHOTO", 1),   # TODO: PASSPORT ENUM 추가 시 교체
    ("프로필", "PROFILE", 1),
    ("개인화보", "PERSONAL_PORTRAIT", 1),
    ("취업", "JOB_PHOTO", 1),
    ("가족", "FAMILY", 2),
    ("우정", "FRIENDSHIP", 2),
]

# 이미지 파일 카테고리 -> 상품명 접두어
IMG_BASE_TO_PRODUCT = {
    "id_photo": "증명", "profile_photo": "프로필", "personal_portrait": "개인화보",
    "job_photo": "취업", "family_photo": "가족", "friendship_photo": "우정",
}

WEEKDAY_KO = {"월": 1, "화": 2, "수": 3, "목": 4, "금": 5, "토": 6, "일": 0}


def category_of(name):
    for prefix, enum, people in CATEGORY:
        if name.startswith(prefix):
            return prefix, enum, people
    raise ValueError(f"카테고리 매핑 실패: {name}")


def parse_products(text):
    """컨셉별가격 문자열을 상품 목록으로 분해한다.

    '개인화보(시크·자연광) ₩148,000' 처럼 괄호 안에 세부 컨셉이 여러 개면
    각각을 별도 상품으로 나눈다. 촬영목적(개인화보)이 중분류, 세부 컨셉이 상품 카드다.
    가격은 원본에 하나뿐이라 동일하게 적용한다.
    """
    products = []
    for part in text.split(" / "):
        m = re.match(r"^(.*?)\s*₩([\d,]+)$", part.strip())
        name = m.group(1)
        price = int(m.group(2).replace(",", ""))

        paren = re.search(r"\((.+)\)$", name)
        subs = (
            [s.strip() for s in paren.group(1).split("·")]
            if paren and "·" in paren.group(1)
            else None
        )

        if subs:
            base = name[: paren.start()]
            prefix, enum, people = category_of(base)
            # subIndex는 이미지 파일명의 번호(개인화보1, 개인화보2...)와 대응한다
            for index, sub in enumerate(subs, start=1):
                products.append({
                    "namePrefix": prefix,
                    "subIndex": index,
                    "name": f"{base}({sub})",
                    "price": price,
                    "shootingCategory": enum,
                    "basePeople": people,
                    "images": [],
                })
        else:
            prefix, enum, people = category_of(name)
            products.append({
                "namePrefix": prefix,
                "subIndex": None,
                "name": name,
                "price": price,
                "shootingCategory": enum,
                "basePeople": people,
                "images": [],
            })
    return products


def parse_hair_makeup(text):
    """'O (A 루아뷰티 +₩58,000 / B 베이지룸 +₩33,000, 사전예약 필수)' -> 파트너 목록"""
    if not str(text).startswith("O"):
        return []
    inner = re.search(r"\((.*)\)$", str(text).strip())
    if not inner:
        return []
    details = []
    for part in inner.group(1).split(" / "):
        m = re.match(r"^([A-Z])\s+(.+?)\s*\+₩([\d,]+)", part.strip())
        if m:
            details.append({
                "partnerName": m.group(2).strip(),
                "additionalPrice": int(m.group(3).replace(",", "")),
                "displayOrder": ord(m.group(1)) - ord("A") + 1,
            })
    return details


def parse_hours(text):
    """운영시간 문자열 -> {weekday:[open,close], weekend:[open,close]}"""
    times = re.findall(r"(\d{2}:\d{2})-(\d{2}:\d{2})", text)
    if "평일" in text and "주말" in text:
        return {"weekday": list(times[0]), "weekend": list(times[1])}
    return {"weekday": list(times[0]), "weekend": list(times[0])}


def parse_closed_weekday(holiday_text, hours_text):
    """정기 휴무 요일 (0=일 ~ 6=토). 없으면 None."""
    m = re.search(r"매주\s*([월화수목금토일])요일", str(holiday_text))
    if m:
        return WEEKDAY_KO[m.group(1)]
    m = re.search(r"\(([월화수목금토일])\s*휴무\)", str(hours_text))
    if m:
        return WEEKDAY_KO[m.group(1)]
    return None


def split_address(addr):
    """'서울 마포구 와우산로 8' -> ('서울 마포구', '와우산로 8')"""
    parts = addr.split(" ", 2)
    if len(parts) >= 3:
        return f"{parts[0]} {parts[1]}", parts[2]
    return addr, ""


def section(*lines):
    return "\n".join([l for l in lines if l])


# ---------- 이미지 매핑 CSV 로드 ----------
images_by_studio = collections.defaultdict(list)
with open(IMG_CSV, encoding="utf-8-sig") as fp:
    for row in csv.DictReader(fp):
        images_by_studio[row["studioExcelId"]].append(row)

# ---------- 엑셀 파싱 ----------
wb = openpyxl.load_workbook(XLSX, data_only=True)
sheet_name = next((s for s in SHEET_CANDIDATES if s in wb.sheetnames), None)
if sheet_name is None:
    sys.exit(f"사진관 시트를 찾을 수 없습니다. 시트 목록: {wb.sheetnames}")
ws = wb[sheet_name]
rows = list(ws.iter_rows(min_row=1, values_only=True))
hdr, data = rows[0], [r for r in rows[1:] if r[0]]
I = {h: i for i, h in enumerate(hdr)}
g = lambda r, k: r[I[k]]

studios, orphan_images = [], []

for r in data:
    excel_id = g(r, "ID")
    products = parse_products(g(r, "컨셉별가격"))

    # --- 이미지를 상품에 연결 ---
    # 파일명 접두어로 중분류를 찾고, 세부 컨셉으로 나뉜 상품은 파일명 끝 번호로 짝을 맞춘다.
    # (개인화보1(시크) -> 개인화보(시크), 개인화보2(자연광) -> 개인화보(자연광))
    for img in images_by_studio.get(excel_id, []):
        parsed = re.match(r"^[A-Z]{2}-\d+_(.+?)(?:_(\d+))?\.jpg$", img["fileName"])
        prefix = IMG_BASE_TO_PRODUCT.get(parsed.group(1)) if parsed else None
        number = int(parsed.group(2)) if parsed and parsed.group(2) else None

        candidates = [p for p in products if p["namePrefix"] == prefix]
        target = None
        if number is not None:
            target = next((p for p in candidates if p["subIndex"] == number), None)
        if target is None:
            # 나뉘지 않은 상품은 이미지를 여러 장 가질 수 있다
            target = next((p for p in candidates if p["subIndex"] is None), None)
        if target is None:
            orphan_images.append((excel_id, img["fileName"]))
            continue
        target["images"].append({
            "url": img["url"],
            "order": len(target["images"]) + 1,
            "studioThumbnailOrder": 1 if img["studioThumbnailOrder"] == "1" else None,
        })

    hair = parse_hair_makeup(g(r, "헤어메이크업연계"))
    services = []
    if hair:
        services.append("HAIR_MAKEUP")
    if str(g(r, "주차")).startswith("O"):
        services.append("PARKING")
    if str(g(r, "의상비치")).startswith("O"):
        services.append("COSTUME")
    if str(g(r, "와이파이")).startswith("O"):
        services.append("WIFI")

    main_addr, sub_addr = split_address(g(r, "주소"))
    lines = [
        LINE_CODE.get(t.strip()) or (int(m.group(1)) if (m := re.match(r"^(\d+)호선$", t.strip())) else None)
        for t in str(g(r, "호선")).split("·")
    ]

    for p in products:
        p.pop("namePrefix")
        p.pop("subIndex")
        p["hasAdditionalPrice"] = bool(hair)

    studios.append({
        "excelId": excel_id,
        "name": g(r, "사진관명"),
        "introduction": g(r, "사진관소개"),
        "notice": g(r, "공지사항"),
        "ratingScore": float(g(r, "별점")),
        "reservationCount": int(g(r, "최근30일예약완료건수")),
        "location": {
            "mainAddress": main_addr,
            "subAddress": sub_addr,
            "locationCategory": LOCATION[g(r, "지역")],
            "latitude": str(g(r, "위도")),
            "longitude": str(g(r, "경도")),
            "nearestStation": g(r, "최근접역"),
            "walkingMinutes": int(g(r, "도보거리(분)")),
            "stationDetail": [c for c in lines if c],
        },
        "products": products,
        "services": services,
        "hairMakeupDetails": hair,
        "schedule": {
            "hours": parse_hours(g(r, "운영시간")),
            "slotMinutes": int(g(r, "슬롯당소요시간(분)")),
            "closedWeekday": parse_closed_weekday(g(r, "휴무일"), g(r, "운영시간")),
        },
        "info": {
            "운영 정보": section(f"영업시간: {g(r, '운영시간')}",
                              f"휴무일: {g(r, '휴무일')}",
                              f"예약 마감: {g(r, '예약마감')}",
                              g(r, "판매자정보")),
            "주차 정보": section(str(g(r, "주차"))),
            "촬영 안내": section(f"의상: {g(r, '의상비치')}",
                              f"와이파이: {g(r, '와이파이')}",
                              g(r, "비고")),
            "환불 안내": section(str(g(r, "취소환불규정"))),
        },
    })

with open(OUT, "w", encoding="utf-8") as fp:
    json.dump(studios, fp, ensure_ascii=False, indent=2)

# ---------- 검증 출력 ----------
n_products = sum(len(s["products"]) for s in studios)
n_images = sum(len(p["images"]) for s in studios for p in s["products"])
n_thumb = sum(1 for s in studios for p in s["products"] for i in p["images"]
              if i["studioThumbnailOrder"] == 1)
print(f"사진관   : {len(studios)}")
print(f"상품     : {n_products}")
print(f"이미지   : {n_images}")
print(f"대표썸네일: {n_thumb}  (사진관당 1개여야 정상)")
print(f"헤어메이크업 연계: {sum(1 for s in studios if s['hairMakeupDetails'])}")
print(f"고아 이미지: {len(orphan_images)} {orphan_images[:5]}")
print(f"\n생성: {OUT}")
